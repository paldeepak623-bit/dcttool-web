"""Sheet data helpers — count rows, detect empty sheets (no forced fill from Membership)."""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

HEADER_ROW = 1

SHEET_LABELS: dict[str, str] = {
  "Declaration Template": "Declaration",
  "Membership": "Membership",
  "ShortTermLoans": "Short Term Loans",
  "LandDetails": "Land Details",
  "LoanTransaction": "Loan Transaction",
  "Savings": "Savings",
  "Dropdown List": "Dropdown List",
  "ShortTermLoans To LoanTransaction": "ShortTermLoans To LoanTransaction",
}


def sheet_label(sheet_name: str) -> str:
  return SHEET_LABELS.get(sheet_name, sheet_name)


def count_data_rows(ws: Worksheet) -> int:
  count = 0
  for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
    if any(v is not None and v != "" for v in row):
      count += 1
  return count


def has_sheet_data(ws: Worksheet) -> bool:
  return count_data_rows(ws) > 0


def normalize_sheet_title(raw_title: str) -> str | None:
  """Normalize various master sheet tab names into canonical internal names."""
  t = raw_title.strip()
  clean = "".join(c for c in t.lower() if c.isalnum())

  # ShortTermLoans To LoanTransaction (must precede STLoans and LoanTxn)
  if ("short" in clean or "st" in clean) and ("loantrans" in clean or "lt" in clean or "to" in clean):
    return "ShortTermLoans To LoanTransaction"

  # Short Term Loans: 'Shorts terms loans', 'Short Term Loans', 'stloans', 'shorttermloan'
  if any(k in clean for k in ["short", "stloan"]):
    if any(w in clean for w in ["term", "loan", "loans"]):
      return "ShortTermLoans"

  # Loan Transaction: 'LoanTransaction', 'Loan Transaction', 'Loan Transactions', 'loantxn'
  if "loantrans" in clean or "loantxn" in clean:
    return "LoanTransaction"

  # Land Details: 'LandDetails', 'Land Details', 'Land Detail', 'land'
  if "land" in clean:
    return "LandDetails"

  # Savings: 'Savings', 'Saving', 'SavingsDetails', 'savings'
  if "saving" in clean:
    return "Savings"

  # Membership: 'Membership', 'Members', 'Member'
  if "member" in clean:
    return "Membership"

  # Declaration: 'Declaration Template', 'Declaration'
  if "declaration" in clean:
    return "Declaration Template"

  # Dropdown: 'Dropdown List', 'Dropdown', 'DropdownList'
  if "dropdown" in clean:
    return "Dropdown List"

  return None


def canonicalize_workbook_sheets(wb: any) -> dict[str, str]:
  """Rename sheets in workbook to canonical names so all downstream code works seamlessly."""
  renamed: dict[str, str] = {}
  for name in list(wb.sheetnames):
    canon = normalize_sheet_title(name)
    if canon and canon != name and canon not in wb.sheetnames:
      wb[name].title = canon
      renamed[name] = canon
  return renamed


def clean_pacs_name(raw_name: str, fallback_stem: str = "") -> str:
  """Clean raw PACS name from Declaration sheet into pure uppercase PACS name.

  Example: 'Dhamora B-Pacs Limited' -> 'DHAMORA'
           'B PACS BHADSAR' -> 'BHADSAR'
           'B-Pacs KHASEPUR (Jarwal Road)' -> 'KHASEPUR'
  """
  import re
  if not raw_name:
    raw_name = fallback_stem

  s = str(raw_name).replace("_", " ").replace("-", " ")
  s = re.sub(r"\(.*?\)", " ", s)
  s = re.sub(r"\.(xlsx|xlsm|xls|pdf)", " ", s, flags=re.IGNORECASE)
  s = re.sub(r"\b[P]?\d{10,}\b", " ", s)

  patterns = [
      r"\bb\s*pacs\b",
      r"\bp\s*pacs\b",
      r"\bm\s*pacs\b",
      r"\bpacs\b",
      r"\blimited\b",
      r"\bltd\.?\b",
      r"\bsss\b",
      r"\bfff\b",
      r"\bcooperative\b",
      r"\bco\s*operative\b",
      r"\bsamiti\b",
      r"\btemplate\b",
      r"\bfinal\b",
      r"\bdct\b",
      r"\bprocessed\b",
  ]
  for p in patterns:
    s = re.sub(p, " ", s, flags=re.IGNORECASE)

  s = re.sub(r"[^a-zA-Z0-9\s]", " ", s)
  s = " ".join(s.split()).strip().upper()

  if not s and fallback_stem:
    return clean_pacs_name(fallback_stem, "")
  return s or "PACS"


def extract_pacs_metadata_from_workbook(wb: any, fallback_stem: str = "") -> tuple[str, str | None, str | None]:
  """Extract (clean_pacs_name, district, pacs_id) from Declaration sheet.

  Returns:
    clean_pacs_name: UPPERCASE clean pacs name (e.g. 'DHAMORA')
    district: e.g. 'RAMPUR' or None
    pacs_id: e.g. '91356302044001' or None
  """
  import re
  ws = None
  for name in wb.sheetnames:
    if "declaration" in name.lower():
      ws = wb[name]
      break

  raw_pacs = None
  district = None
  pacs_id = None

  if ws is not None:
    for r in range(1, min(ws.max_row + 1, 35)):
      for c in range(1, min(ws.max_column + 1, 10)):
        val = ws.cell(r, c).value
        if not val:
          continue
        s_val = str(val).strip().lower()

        # Check District
        if "district" in s_val and not district:
          for c_next in range(c + 1, min(ws.max_column + 1, 10)):
            nxt = ws.cell(r, c_next).value
            if nxt and str(nxt).strip() and "name" not in str(nxt).lower():
              district = str(nxt).strip().upper()
              break

        # Check PACS Name
        if ("pacs name" in s_val or s_val == "pacs") and not raw_pacs:
          for c_next in range(c + 1, min(ws.max_column + 1, 10)):
            nxt = ws.cell(r, c_next).value
            if nxt and str(nxt).strip() and not str(nxt).strip().isdigit() and "name" not in str(nxt).lower():
              raw_pacs = str(nxt).strip()
              break

        # Check PACS ID
        if ("pacs id" in s_val or s_val == "pacs_id" or "pacs code" in s_val) and not pacs_id:
          for c_next in range(c + 1, min(ws.max_column + 1, 10)):
            nxt = ws.cell(r, c_next).value
            if nxt and str(nxt).strip() and "name" not in str(nxt).lower():
              raw_id = str(nxt).strip()
              m = re.search(r"(\d{10,})", raw_id)
              if m:
                pacs_id = raw_id
                break
              elif raw_id.isalnum() and len(raw_id) >= 8:
                pacs_id = raw_id
                break

  clean_name = clean_pacs_name(raw_pacs or fallback_stem, fallback_stem)
  return clean_name, district, pacs_id

