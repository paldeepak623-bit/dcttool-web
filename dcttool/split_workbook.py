"""Split processed master workbook sheets into individual template files."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from dcttool.date_utils import excel_date_value, export_date_value, is_date_column, parse_date
from dcttool.processors.cross_sheet import to_float
from dcttool.processors.membership_processor import resolve_age
from dcttool.template_files import HEADER_ALIASES, resolve_upload_template
from dcttool.text_utils import sanitize_cell_value

EXPORT_MAP: dict[str, tuple[str, str]] = {
    "Declaration Template": ("DeclarationTemplate.xlsx", "Declaration Template"),
    "Declaration": ("DeclarationTemplate.xlsx", "Declaration Template"),
    "Membership": ("MembershipTemplate.xlsx", "Membership"),
    "ShortTermLoans": ("STLoansTemplate.xlsx", "ShortTermLoans"),
    "Shorts terms loans": ("STLoansTemplate.xlsx", "ShortTermLoans"),
    "Short Term Loans": ("STLoansTemplate.xlsx", "ShortTermLoans"),
    "LandDetails": ("LandDetailsTemplate.xlsx", "LandDetails"),
    "Land Details": ("LandDetailsTemplate.xlsx", "LandDetails"),
    "LoanTransaction": ("LoanTransactionTemplate.xlsx", "LoanTransaction"),
    "Loan Transaction": ("LoanTransactionTemplate.xlsx", "LoanTransaction"),
    "Savings": ("SavingsTemplate.xlsx", "Savings"),
    "Dropdown List": ("DropdownList.xlsx", "Dropdown List"),
    "ShortTermLoans To LoanTransaction": ("ShortTermLoans_To_LoanTransaction.xlsx", "LoanTransaction"),
}

HEADER_ROW = 1
DATA_START_ROW = 2


def _header_col_map(ws: Worksheet, header_row: int = HEADER_ROW) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(header_row, col).value
        if raw is None:
            continue
        name = str(raw).strip()
        if name:
            headers[name] = col
    return headers


def _source_col_for_template_header(
    tpl_header: str,
    src_headers: dict[str, int],
) -> int | None:
    if tpl_header in src_headers:
        return src_headers[tpl_header]
    for alias in HEADER_ALIASES.get(tpl_header, []):
        if alias in src_headers:
            return src_headers[alias]
    # Fuzzy alnum check
    tpl_clean = "".join(c for c in tpl_header.lower() if c.isalnum())
    for s_hdr, col in src_headers.items():
        s_clean = "".join(c for c in s_hdr.lower() if c.isalnum())
        if s_clean == tpl_clean:
            return col
    # Check aliases alnum
    for alias in HEADER_ALIASES.get(tpl_header, []):
        a_clean = "".join(c for c in alias.lower() if c.isalnum())
        for s_hdr, col in src_headers.items():
            s_clean = "".join(c for c in s_hdr.lower() if c.isalnum())
            if s_clean == a_clean:
                return col
    return None


def _clear_template_data_rows(target: Worksheet, start_row: int = DATA_START_ROW) -> None:
    max_row = target.max_row or start_row
    max_col = target.max_column or 1
    if max_row < start_row:
        return
    for row in range(start_row, max_row + 1):
        for col in range(1, max_col + 1):
            cell = target.cell(row, col)
            if cell.value is not None:
                cell.value = None


def _source_data_rows(source: Worksheet, start_row: int = DATA_START_ROW) -> list[int]:
    rows: list[int] = []
    for offset, row in enumerate(
        source.iter_rows(min_row=start_row, values_only=True),
        start=start_row,
    ):
        if any(v is not None and v != "" for v in row):
            rows.append(offset)
    return rows


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _col_header_map(ws: Worksheet, header_row: int = HEADER_ROW) -> dict[int, str]:
    headers: dict[int, str] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(header_row, col).value
        if raw is None:
            continue
        name = str(raw).strip()
        if name:
            headers[col] = name
    return headers


def _export_value(cell_value: Any, *, header_name: str | None = None) -> tuple[Any, str | None]:
    """Strip symbols; format dates; preserve formulas and IDs."""
    if cell_value is None:
        return None, None

    if _is_formula(cell_value):
        return str(cell_value), None

    if header_name in ("PacsIDPKey", "PacsIDPkey", "PacsID", "DCCBSBACNO", "AdhaarCardNo", "AadhaarCardNo"):
        return str(cell_value).strip(), "@"

    if header_name in ("ShareBalance", "ThriftBalance", "DividentBalance"):
        try:
            return round(float(cell_value), 2), "0.00"
        except (ValueError, TypeError):
            return cell_value, None

    if header_name and is_date_column(header_name):
        date_value, date_fmt = export_date_value(cell_value)
        if header_name in ("AdmissionDate", "SanctionedDate", "TransactionDate", "DepositDate"):
            parsed_check = parse_date(date_value)
            if parsed_check is not None and parsed_check >= date(2025, 3, 30):
                date_value = excel_date_value(date(2025, 3, 30))
                date_fmt = "dd-mm-yyyy"
        if date_fmt:
            return date_value, date_fmt
        if date_value is not None:
            return date_value, None

    original = cell_value
    result = sanitize_cell_value(cell_value)
    if result is None and original is not None and str(original).strip() != "":
        return original, None
    return result, None


def export_sheet_into_upload_template(
    source: Worksheet,
    template_path: Path,
    output_path: Path,
    *,
    pacs_id: str | None = None,
) -> None:
    """Paste master values into bundled portal template (headers unchanged, values only)."""
    tpl_wb = openpyxl.load_workbook(template_path, read_only=False, keep_vba=False)
    target = tpl_wb.active

    tpl_headers = _header_col_map(target)
    src_headers = _header_col_map(source)
    _clear_template_data_rows(target)

    data_rows = _source_data_rows(source)
    # Pre-compute column mappings once to avoid repeated lookups
    col_mapping = []
    for tpl_hdr, tpl_col in tpl_headers.items():
        src_col = _source_col_for_template_header(tpl_hdr, src_headers)
        if src_col is None:
            # Positional fallbacks for critical type columns
            if tpl_hdr == "CustomerTypeDescription" and source.max_column >= 1:
                src_col = 1
            elif tpl_hdr == "MemberTypeDescription" and source.max_column >= 2:
                src_col = 2
            elif tpl_hdr == "FarmerTypeDescription":
                for sh, sc in src_headers.items():
                    if "farmer" in sh.lower():
                        src_col = sc
                        break
        # Include type columns even if src_col is None so defaults will be applied
        if src_col is not None or tpl_hdr in (
            "CustomerTypeDescription", "MemberTypeDescription", "FarmerTypeDescription",
            "BranchId", "PacsIDPKey", "PacsID", "PacsCode", "VillageDescription"
        ):
            col_mapping.append((tpl_hdr, tpl_col, src_col))

    dob_col = src_headers.get("DOB")
    adm_col = src_headers.get("AdmissionDate")

    for out_row, src_row in enumerate(data_rows, start=DATA_START_ROW):
        share_val = 0.0
        if "ShareBalance" in src_headers:
            share_val = to_float(source.cell(src_row, src_headers["ShareBalance"]).value) or 0.0

        for tpl_hdr, tpl_col, src_col in col_mapping:
            src_cell = source.cell(src_row, src_col) if src_col else None
            src_val = src_cell.value if src_cell else None

            # Special Type Handling: Ensure Type columns are NEVER left blank
            if tpl_hdr == "CustomerTypeDescription":
                raw = str(src_val or "").strip()
                if "NOMINAL" in raw.upper():
                    val = "Nominal Member"
                elif "MEMBER" in raw.upper():
                    val = "Member"
                elif raw and raw.lower() not in ("none", "null"):
                    val = raw.title()
                else:
                    val = "Member" if share_val >= 100 else "Nominal Member"
                target.cell(out_row, tpl_col, value=val)
                continue

            if tpl_hdr == "MemberTypeDescription":
                raw = str(src_val or "").strip()
                if "B" in raw.upper() and ("TYPE" in raw.upper() or raw.strip().upper() == "B"):
                    val = "B Type"
                elif "A" in raw.upper() and ("TYPE" in raw.upper() or raw.strip().upper() == "A"):
                    val = "A Type"
                elif raw and raw.lower() not in ("none", "null"):
                    val = raw.title()
                else:
                    val = "A Type" if share_val >= 100 else "B Type"
                target.cell(out_row, tpl_col, value=val)
                continue

            if tpl_hdr == "FarmerTypeDescription":
                raw = str(src_val or "").strip()
                if "BIG" in raw.upper():
                    val = "Big"
                elif "MED" in raw.upper():
                    val = "Medium"
                elif "SMALL" in raw.upper():
                    val = "Medium"
                elif raw and raw.title() in ("Big", "Medium"):
                    val = raw.title()
                else:
                    val = "Medium"
                target.cell(out_row, tpl_col, value=val)
                continue

            # Special Aadhaar Handling: Ensure Aadhaar card numbers are completely cleared (blank)
            if any(k in tpl_hdr.lower() for k in ("adhaar", "aadhar", "adhar")):
                target.cell(out_row, tpl_col, value=None)
                continue

            # Special Address Handling: Ensure Address1 and Address2 are completely cleared (blank)
            if tpl_hdr in ("Address1", "Address2", "Address", "Address 1", "Address 2"):
                target.cell(out_row, tpl_col, value=None)
                continue

            # Special BranchId Handling: Ensure BranchId is ALWAYS 101 across all templates
            if tpl_hdr in ("BranchId", "BranchID", "Branch Id", "Branch_Id"):
                dest = target.cell(out_row, tpl_col, value=101)
                dest.number_format = "0"
                continue

            # Special PacsID Handling: Ensure PacsID / PacsIDPKey is NEVER blank across all templates
            if tpl_hdr in ("PacsIDPKey", "PacsID", "PacsId", "PacsCode", "PacsIDPkey"):
                raw_pacs = str(src_val or "").strip()
                if not raw_pacs or raw_pacs.lower() in ("none", "null", ""):
                    final_pacs = str(pacs_id or "").strip()
                    if not final_pacs:
                        from dcttool.sheet_populator import extract_pacs_metadata_from_workbook
                        _, _, decl_pacs = extract_pacs_metadata_from_workbook(source.parent, fallback_stem="")
                        final_pacs = decl_pacs or "91351905054001"
                else:
                    final_pacs = raw_pacs
                dest = target.cell(out_row, tpl_col, value=str(final_pacs))
                dest.number_format = "@"
                continue

            # Special Village Handling: Ensure VillageDescription is NEVER left blank or Others
            if tpl_hdr in ("VillageDescription", "Village", "VillageName", "Village Description"):
                raw_v = str(src_val or "").strip()
                if not raw_v or raw_v.lower() in ("none", "null", "others", "--select--", ""):
                    fallback_v = "Main Village"
                    for r_chk in range(DATA_START_ROW, min(source.max_row + 1, 100)):
                        cv = source.cell(r_chk, src_col).value if src_col else None
                        if cv and str(cv).strip().lower() not in ("none", "null", "others", "--select--", ""):
                            fallback_v = str(cv).strip()
                            break
                    val = fallback_v
                else:
                    val = raw_v
                target.cell(out_row, tpl_col, value=val)
                continue

            # Special LandDetails AdmissionNo: Must strictly be 8-digit zero-padded text (e.g. 00000012)
            if target.title in ("LandDetails", "Land Details") and tpl_hdr in ("AdmissionNo", "Admission No", "Admission_No"):
                raw_adm = str(src_val or "").strip()
                if raw_adm and raw_adm.isdigit():
                    val = raw_adm.zfill(8)
                else:
                    val = raw_adm
                dest = target.cell(out_row, tpl_col, value=val)
                dest.number_format = "@"
                continue

            # Special LandDetails StateName: Always "UTTAR PRADESH"
            if target.title in ("LandDetails", "Land Details") and tpl_hdr in ("StateName", "State Name"):
                target.cell(out_row, tpl_col, value="UTTAR PRADESH")
                continue

            # Special Age Handling: Ensure integer age is always populated
            if tpl_hdr == "Age":
                raw_age = src_cell.value if src_cell else None
                dob_val = parse_date(source.cell(src_row, dob_col).value) if dob_col else None
                adm_val = parse_date(source.cell(src_row, adm_col).value) if adm_col else None
                age_val = resolve_age(raw_age, dob_val, adm_val)
                dest = target.cell(out_row, tpl_col, value=int(age_val))
                dest.number_format = "0"
                continue

            # Special Balance Handling: Ensure numeric values are always preserved and formatted
            if tpl_hdr in ("ShareBalance", "ThriftBalance", "DividentBalance"):
                raw_bal = src_cell.value if src_cell else None
                if raw_bal is not None and str(raw_bal).strip() != "":
                    try:
                        dest = target.cell(out_row, tpl_col, value=round(float(raw_bal), 2))
                        dest.number_format = "0.00"
                    except (ValueError, TypeError):
                        dest = target.cell(out_row, tpl_col, value=raw_bal)
                continue

            # Special Value Handling: LoanTransaction CollectedPrincipal must be numeric value
            if tpl_hdr == "CollectedPrincipal":
                val = src_cell.value if src_cell else 0.0
                try:
                    dest = target.cell(out_row, tpl_col, value=round(float(val), 2))
                    dest.number_format = "0.00"
                except (ValueError, TypeError):
                    dest = target.cell(out_row, tpl_col, value=0.0)
                    dest.number_format = "0.00"
                continue

            if src_cell is None:
                continue

            value, date_fmt = _export_value(src_cell.value, header_name=tpl_hdr)
            if value is None:
                continue
            dest = target.cell(out_row, tpl_col, value=value)
            if date_fmt:
                dest.number_format = date_fmt
            elif src_cell.number_format:
                dest.number_format = src_cell.number_format

    output_path.parent.mkdir(parents=True, exist_ok=True)
    tpl_wb.save(output_path)
    tpl_wb.close()


def export_sheet(
    workbook: Workbook,
    sheet_name: str,
    output_path: Path,
    *,
    tab_name: str | None = None,
) -> None:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
    source = workbook[sheet_name]
    out = Workbook()
    target = out.active
    target.title = (tab_name or sheet_name)[:31]

    col_headers = _col_header_map(source)
    for row in range(1, (source.max_row or 1) + 1):
        for col in range(1, (source.max_column or 1) + 1):
            src_cell = source.cell(row, col)
            if src_cell.value is None:
                continue
            header_name = col_headers.get(col) if row >= DATA_START_ROW else None
            if header_name and header_name in ("Address1", "Address2", "Address", "Address 1", "Address 2"):
                continue
            value, date_fmt = _export_value(src_cell.value, header_name=header_name)
            if value is None:
                continue
            dest = target.cell(row, col, value=value)
            if date_fmt:
                dest.number_format = date_fmt
            elif src_cell.number_format:
                dest.number_format = src_cell.number_format

    for key, dim in source.column_dimensions.items():
        if dim.width:
            target.column_dimensions[key].width = dim.width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    out.save(output_path)
    out.close()


def find_source_sheet(workbook: Workbook, sheet_name: str) -> Worksheet | None:
    """Find sheet in workbook allowing for name aliases and casing differences."""
    if sheet_name == "ShortTermLoans To LoanTransaction":
        if "ShortTermLoans_To_LT" in workbook.sheetnames:
            return workbook["ShortTermLoans_To_LT"]
        if "ShortTermLoans To LoanTransaction" in workbook.sheetnames:
            return workbook["ShortTermLoans To LoanTransaction"]
        if "LoanTransaction" in workbook.sheetnames:
            return workbook["LoanTransaction"]

    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]

    # Clean alphanumeric comparison
    target_clean = "".join(c for c in sheet_name.lower() if c.isalnum())
    for s in workbook.sheetnames:
        if "".join(c for c in s.lower() if c.isalnum()) == target_clean:
            return workbook[s]

    # Normalize aliases from sheet_populator
    from dcttool.sheet_populator import normalize_sheet_title
    for s in workbook.sheetnames:
        if normalize_sheet_title(s) == sheet_name:
            return workbook[s]

    return None


def export_selected_templates(
    workbook: Workbook,
    output_dir: Path,
    sheet_names: list[str],
    *,
    pacs_id: str | None = None,
) -> list[Path]:
    from dcttool.sheet_populator import normalize_sheet_title
    if not pacs_id:
        try:
            from dcttool.sheet_populator import extract_pacs_metadata_from_workbook
            _, _, decl_pacs = extract_pacs_metadata_from_workbook(workbook, fallback_stem="")
            pacs_id = decl_pacs or "91351905054001"
        except Exception:
            pacs_id = "91351905054001"

    saved: list[Path] = []
    for sheet_name in sheet_names:
        canon = normalize_sheet_title(sheet_name) or sheet_name
        if canon in EXPORT_MAP:
            filename, tab = EXPORT_MAP[canon]
        elif sheet_name in EXPORT_MAP:
            filename, tab = EXPORT_MAP[sheet_name]
        else:
            filename = f"{sheet_name.replace(' ', '')}Template.xlsx"
            tab = sheet_name
        path = output_dir / filename
        upload_tpl = resolve_upload_template(canon) or resolve_upload_template(sheet_name)

        src_sheet = find_source_sheet(workbook, sheet_name) or find_source_sheet(workbook, canon)
        if src_sheet is None:
            continue

        if upload_tpl is not None:
            export_sheet_into_upload_template(src_sheet, upload_tpl, path, pacs_id=pacs_id)
        else:
            export_sheet(workbook, src_sheet.title, path, tab_name=tab)
        saved.append(path)
        if sheet_name in ("ShortTermLoans To LoanTransaction", "ShortTermLoans_To_LT") or canon == "ShortTermLoans To LoanTransaction":
            alt_path = output_dir / "ShortTermLoans To LoanTransaction.xlsx"
            if path != alt_path and path.exists():
                try:
                    shutil.copy2(path, alt_path)
                    saved.append(alt_path)
                except Exception:
                    pass
    return saved


def save_full_master_copy(
    workbook: Workbook,
    output_dir: Path,
    original_path: Path,
) -> Path:
    """Save processed master workbook directly named after PACS Name from Declaration sheet."""
    from dcttool.sheet_populator import extract_pacs_metadata_from_workbook
    clean_pacs, _, _ = extract_pacs_metadata_from_workbook(workbook, fallback_stem=original_path.stem)

    # Strictly name the Excel sheet by PACS Name, NOT PACS ID
    base_name = clean_pacs if clean_pacs and clean_pacs != "PACS" else original_path.stem
    out_path = output_dir / f"{base_name}_DCT_Processed.xlsx"
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)

    # Also save a direct PACS Name copy (e.g. PAHAREMAU.xlsx)
    direct_pacs_path = output_dir / f"{base_name}.xlsx"
    if direct_pacs_path != out_path:
        try:
            import shutil
            shutil.copy2(out_path, direct_pacs_path)
        except Exception:
            pass

    return out_path
