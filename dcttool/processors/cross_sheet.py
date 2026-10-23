"""Cross-template helpers: admission keys, crop normalize, thrift transfers."""

from __future__ import annotations

from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

HEADER_ROW = 1
ZERO_FIX_AMOUNT = 1000.0


def normalize_admission(value: Any) -> str | None:
  if value is None:
    return None
  if isinstance(value, float) and value.is_integer():
    return str(int(value))
  text = str(value).strip()
  if not text or text.upper() in {"NA", "N/A", "NONE", "NULL"}:
    return None
  if text.endswith(".0"):
    text = text[:-2]
  return text.lstrip("0") or "0"


def to_float(value: Any) -> float | None:
  if value is None:
    return None
  if isinstance(value, (int, float)):
    return float(value)
  text = str(value).strip().replace(",", "")
  if not text:
    return None
  try:
    return float(text)
  except ValueError:
    return None


def header_map(ws: Worksheet) -> dict[str, int]:
  headers: dict[str, int] = {}
  for col in range(1, ws.max_column + 1):
    raw = ws.cell(HEADER_ROW, col).value
    if raw is None:
      continue
    name = str(raw).strip()
    if name:
      headers[name] = col
  return headers


def collect_loan_admissions(*sheets: Worksheet | None) -> set[str]:
  """AdmissionNos that appear on ShortTermLoans, LoanTransaction, LandDetails, Savings, etc."""
  found: set[str] = set()
  for ws in sheets:
    if ws is None:
      continue
    headers = header_map(ws)
    adm_cols = [
      col for name, col in headers.items()
      if "admission" in name.lower() and "gpa" not in name.lower() and "date" not in name.lower()
    ]
    if not adm_cols:
      continue
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
      for col in adm_cols:
        key = normalize_admission(ws.cell(row, col).value)
        if key:
          found.add(key)
  return found


def normalize_crop_description(raw: Any) -> str:
  """
  As-per-sheet crop mapping:
  Animal Husbandry → Animals
  Wheat → Wheat
  Paddy / blank → Paddy
  """
  if raw is None or str(raw).strip() == "":
    return "Paddy"
  text = " ".join(str(raw).strip().split())
  low = text.lower()
  if "animal" in low:
    return "Animals"
  if "wheat" in low:
    return "Wheat"
  if "paddy" in low:
    return "Paddy"
  # Title-case cleaned value for other crops if any
  return text.title()


def build_st_crop_by_loan(ws: Worksheet | None) -> dict[str, str]:
  """Map LoanNo / AdmissionNo → normalized CropDescription from ShortTermLoans."""
  out: dict[str, str] = {}
  if ws is None:
    return out
  headers = header_map(ws)
  adm_col = headers.get("AdmissionNo")
  loan_col = headers.get("LoanNo")
  crop_col = headers.get("CropDescription")
  if not crop_col:
    return out
  for row in range(HEADER_ROW + 1, ws.max_row + 1):
    crop = normalize_crop_description(ws.cell(row, crop_col).value)
    if loan_col:
      loan_key = normalize_admission(ws.cell(row, loan_col).value)
      if loan_key:
        out[f"loan:{loan_key}"] = crop
    if adm_col:
      adm_key = normalize_admission(ws.cell(row, adm_col).value)
      if adm_key:
        out.setdefault(f"adm:{adm_key}", crop)
  return out


def abs_negative_pair(principal: float | None, interest: float | None) -> float | None:
  """If both values are negative, return abs(p)+abs(i); else None."""
  if principal is None or interest is None:
    return None
  if principal < 0 and interest < 0:
    return round(abs(principal) + abs(interest), 2)
  return None
